"""
The deck design system: the console's crimson theme (src/prama/web/static/css/themes.css,
the light theme), so the deck and the product look like one thing. Crimson for accents and
table heads, the deep heading crimson for titles, and the console header's gradient (deep
crimson through crimson to indigo) for the dark slides. The prism's six-colour spectrum is
the one flourish (docs/reference/brand.md). Calibri throughout, bold for headings.

Every primitive that places text measures it with ``metrics`` — the same
estimator the geometry audit uses — and the fitting helpers shrink the type or
refuse, so a slide that would overflow fails the build rather than the reader.

Copyright (c) 2026 Ashutosh Sinha. All rights reserved.
Proprietary and confidential. See LICENSE and NOTICE at the repository root.
"""

from __future__ import annotations

from collections.abc import Callable
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt
from pptx.util import Inches as In

from metrics import FOOTER_Y, SAFETY, SANS, SH, SW, est_lines, text_extent, text_h

# The console's light theme, token for token (--maya-*). The names below are the deck's
# roles; REFRACT and INDIGO keep their old names so every layout follows the new palette.
REFRACT = RGBColor(0xA5, 0x1C, 0x30)  # --maya-crimson: accents, table heads, numbers
HEADING = RGBColor(0x6E, 0x11, 0x20)  # --maya-heading: slide titles
INDIGO = RGBColor(0x5C, 0x0E, 0x1B)  # --maya-nav-from: the dark surfaces
NAV_TO = RGBColor(0x29, 0x33, 0x52)  # --maya-nav-to: where the dark gradient ends
HAZE = RGBColor(0xF3, 0xC6, 0xCF)  # --maya-glow: labels on the dark surfaces
MIST = RGBColor(0xFB, 0xEE, 0xF0)  # --maya-crimson-tint: text on dark; tinted rows
TINT = RGBColor(0xFB, 0xEE, 0xF0)  # --maya-crimson-tint
INK = RGBColor(0x1A, 0x1A, 0x1A)  # --maya-ink
SLATE = RGBColor(0x63, 0x6B, 0x76)  # --maya-slate
MUTED = RGBColor(0x8E, 0x94, 0x9C)
RULE = RGBColor(0xE3, 0xDE, 0xD7)  # --maya-border
FROST = RGBColor(0xF7, 0xF5, 0xF2)  # --maya-canvas: alternate rows, tiles
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOLD = RGBColor(0xE8, 0xB3, 0x3A)
OK = RGBColor(0x1E, 0x6B, 0x3A)  # --maya-ok
WARN = RGBColor(0x7A, 0x52, 0x00)  # --maya-warn
#: The prism's six rays, one per quality dimension (docs/reference/brand.md). On a slide
#: they are always the six together, in this order, never one used decoratively.
SPECTRUM = [
    RGBColor(0x00, 0xB3, 0xA4),  # accuracy
    RGBColor(0x2F, 0xB6, 0x73),  # completeness
    RGBColor(0x8C, 0xBF, 0x3F),  # consistency
    RGBColor(0xE8, 0xB3, 0x3A),  # timeliness
    RGBColor(0xE8, 0x82, 0x3A),  # uniqueness
    RGBColor(0xD9, 0x53, 0x4F),  # validity
]
NAVY = RGBColor(0x2E, 0x40, 0x57)

ML = 0.85
CW = SW - 2 * ML
BODY_BOTTOM = FOOTER_Y - 0.12

_state: dict[str, Any] = {"prs": None, "chapter": "", "n": 0}


class DoesNotFit(Exception):
    """A block that cannot be made to fit its box at the smallest permitted size."""


def new_deck(chapter: str) -> Any:
    """Start a fresh 16:9 presentation and make it the current one."""
    prs = Presentation()
    prs.slide_width = In(SW)
    prs.slide_height = In(SH)
    _state.update(prs=prs, chapter=chapter, n=0)
    return prs


def chapter(name: str) -> None:
    _state["chapter"] = name


def blank() -> Any:
    prs = _state["prs"]
    return prs.slides.add_slide(prs.slide_layouts[6])


def remove(shape: Any) -> None:
    el = shape._element
    el.getparent().remove(el)


def rect(
    sl: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: Any = None,
    line: Any = None,
    lw: float = 1.0,
    shape: Any = MSO_SHAPE.RECTANGLE,
) -> Any:
    s = sl.shapes.add_shape(shape, In(x), In(y), In(w), In(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    return s


def txt(
    sl: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    align: Any = PP_ALIGN.LEFT,
    anchor: Any = MSO_ANCHOR.TOP,
) -> Any:
    tb = sl.shapes.add_textbox(In(x), In(y), In(w), In(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return tf


def para(
    tf: Any,
    text: str,
    size: float = 14,
    color: Any = INK,
    bold: bool = False,
    font: str = SANS,
    italic: bool = False,
    space_after: float = 6,
    first: bool = False,
    line: float = 1.2,
    space_before: float = 0,
) -> Any:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return p


def runs(
    tf: Any,
    parts: list[tuple],
    size: float = 14,
    space_after: float = 6,
    first: bool = False,
    line: float = 1.2,
    level: int = 0,
    space_before: float = 0,
) -> Any:
    """``parts`` is a list of ``(text, color, bold)`` tuples, one run each."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    p.level = level
    for text, color, bold in parts:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = SANS
    return p


def fitted(
    sl: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    write: Callable[[Any, float], None],
    start: float,
    floor: float = 9.0,
) -> float:
    """Write a text block at the largest size in [floor, start] whose estimated
    extent fits ``h``; return the height it uses. Raises ``DoesNotFit``."""
    size = start
    while size >= floor:
        tf = txt(sl, x, y, w, h)
        write(tf, size)
        need = text_extent(sl.shapes[-1])
        if need <= h:
            return need
        remove(sl.shapes[-1])
        size -= 0.5
    raise DoesNotFit(f"text block does not fit {w:.2f}x{h:.2f} at {floor}pt")


def spectrum(sl: Any, x: float, y: float, w: float, h: float, vertical: bool = False) -> None:
    """The prism's six rays as one band, split evenly along its length."""
    n = len(SPECTRUM)
    for i, colour in enumerate(SPECTRUM):
        if vertical:
            rect(sl, x, y + i * h / n, w, h / n, fill=colour)
        else:
            rect(sl, x + i * w / n, y, w / n, h, fill=colour)


def ground(sl: Any) -> Any:
    """The dark slides' ground: the console header's gradient, deep crimson through
    crimson to indigo, left to right."""
    s = rect(sl, 0, 0, SW, SH, fill=INDIGO)
    s.fill.gradient()
    s.fill.gradient_angle = 0
    stops = s.fill.gradient_stops
    stops[0].color.rgb, stops[0].position = INDIGO, 0.0
    stops[1].color.rgb, stops[1].position = NAV_TO, 1.0
    # A third stop, the console's crimson, a little past the middle: python-pptx makes two.
    from copy import deepcopy

    mid = deepcopy(stops._gsLst[0])
    mid.set("pos", "55000")
    mid[0].set("val", "A51C30")
    stops._gsLst.insert(1, mid)
    return s


def circle(
    sl: Any, x: float, y: float, d: float, label: str, fill: Any = None, size: float = 14
) -> Any:
    """A filled circle with a centred label: the deck's numbered-step motif."""
    s = rect(sl, x, y, d, d, fill=fill or REFRACT, shape=MSO_SHAPE.OVAL)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = SANS
    return s


def footer(sl: Any) -> None:
    rect(sl, ML, FOOTER_Y, CW, 0.008, fill=RULE)
    tf = txt(sl, ML, SH - 0.46, CW * 0.7, 0.24)
    label = "PRAMA  •  Ashutosh Sinha"
    if _state["chapter"]:
        label += f"  •  {_state['chapter']}"
    para(tf, label, size=8.5, color=MUTED, first=True, space_after=0)
    tf = txt(sl, ML + CW * 0.7, SH - 0.46, CW * 0.3, 0.24, align=PP_ALIGN.RIGHT)
    para(tf, str(_state["n"]), size=8.5, color=MUTED, first=True, space_after=0, bold=True)


def content(title: str, kicker: str | None = None) -> tuple[Any, float]:
    """A content slide's chrome. Returns ``(slide, body_top)``."""
    _state["n"] += 1
    sl = blank()
    rect(sl, 0, 0, SW, SH, fill=WHITE)
    y = 0.48
    if kicker:
        tf = txt(sl, ML, y, CW, 0.24)
        para(tf, kicker.upper(), size=10.5, color=REFRACT, bold=True, first=True, space_after=0)
        y += 0.30
    size = 28.0
    while size > 20 and est_lines(title, CW * SAFETY, size, True, SANS) > 1:
        size -= 1
    th = text_h(title, CW * SAFETY, size, True, SANS, 1.1)
    tf = txt(sl, ML, y, CW, th + 0.04)
    para(tf, title, size=size, color=HEADING, bold=True, first=True, space_after=0, line=1.1)
    body_top = y + th + 0.26
    rect(sl, ML, body_top - 0.14, CW, 0.012, fill=RULE)
    footer(sl)
    return sl, body_top


def divider(num: str, title: str, sub: str, points: list[str]) -> Any:
    """A section divider on the dark gradient: a large numbered circle, the section's
    title and its promise, and what it covers on the right."""
    _state["chapter"] = title
    _state["n"] += 1
    sl = blank()
    ground(sl)
    d = 1.45
    circle(sl, ML + 0.1, 1.15, d, num, fill=REFRACT, size=48)
    tw = CW * 0.56
    size = 40.0
    while size > 26 and est_lines(title, tw * SAFETY, size, True, SANS) > 2:
        size -= 2
    th = text_h(title, tw * SAFETY, size, True, SANS, 1.05)
    tf = txt(sl, ML + 0.1, 2.95, tw, th + 0.05)
    para(tf, title, size=size, color=WHITE, bold=True, first=True, space_after=0, line=1.05)
    top = 2.95 + th + 0.30
    fitted(
        sl,
        ML + 0.1,
        top,
        tw,
        SH - 0.7 - top,
        lambda tf, s: para(
            tf, sub, size=s, color=HAZE, italic=True, first=True, space_after=0, line=1.3
        ),
        17,
        11,
    )
    x = ML + CW * 0.64
    tf = txt(sl, x, 1.2, CW * 0.36, 0.3)
    para(tf, "IN THIS SECTION", size=10, color=HAZE, bold=True, first=True, space_after=0)

    def write(tf: Any, s: float) -> None:
        for i, pnt in enumerate(points):
            runs(
                tf,
                [(f"{i + 1:02d}   ", HAZE, True), (pnt, WHITE, False)],
                size=s,
                space_after=9,
                first=i == 0,
                line=1.15,
            )

    fitted(sl, x, 1.62, CW * 0.36, 4.6, write, 15, 10)
    return sl


def _row_heights(
    data: list[list[str]], widths: list[float], fs: float, hfs: float, bold_col0: bool
) -> list[float]:
    heights = []
    for r, row in enumerate(data):
        size = hfs if r == 0 else fs
        need = 0.0
        for c, cell in enumerate(row):
            bold = r == 0 or (bold_col0 and c == 0)
            need = max(need, text_h(cell, (widths[c] - 0.18) * SAFETY, size, bold, SANS, 1.0))
        heights.append(max(0.30, need + 0.12))
    return heights


def _fill_cell(
    cell: Any, text: str, r: int, c: int, fs: float, hfs: float, bold_col0: bool
) -> None:
    cell.margin_left = In(0.09)
    cell.margin_right = In(0.07)
    cell.margin_top = cell.margin_bottom = In(0.035)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.fill.solid()
    cell.fill.fore_color.rgb = REFRACT if r == 0 else (FROST if r % 2 == 0 else WHITE)
    tf = cell.text_frame
    tf.word_wrap = True
    run = tf.paragraphs[0].add_run()
    run.text = str(text)
    run.font.name = SANS
    run.font.size = Pt(hfs if r == 0 else fs)
    run.font.bold = r == 0 or (bold_col0 and c == 0)
    run.font.color.rgb = WHITE if r == 0 else (INDIGO if bold_col0 and c == 0 else INK)


def table(
    sl: Any,
    data: list[list[str]],
    x: float,
    y: float,
    w: float,
    col_w: list[float] | None = None,
    fs: float = 11.5,
    hfs: float = 11,
    bold_col0: bool = True,
) -> float:
    """A table whose row heights are computed from its wrapped text.
    Returns the rendered height, so callers place what follows from it."""
    cols = len(data[0])
    col_w = col_w or [1.0] * cols
    widths = [w * c / sum(col_w) for c in col_w]
    heights = _row_heights(data, widths, fs, hfs, bold_col0)
    total = sum(heights)
    gf = sl.shapes.add_table(len(data), cols, In(x), In(y), In(w), In(total))
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    for i, cw in enumerate(widths):
        tbl.columns[i].width = Emu(int(In(cw)))
    for r, row in enumerate(data):
        tbl.rows[r].height = Emu(int(In(heights[r])))
        for c, cell in enumerate(row):
            _fill_cell(tbl.cell(r, c), cell, r, c, fs, hfs, bold_col0)
    return total


def fitted_table(
    sl: Any,
    data: list[list[str]],
    x: float,
    y: float,
    w: float,
    h: float,
    col_w: list[float] | None = None,
    start: float = 14.0,
    bold_col0: bool = True,
) -> float:
    """``table`` at the largest size in [9, start] that fits ``h``."""
    fs = start
    while fs >= 9.0:
        used = table(sl, data, x, y, w, col_w, fs=fs, hfs=min(fs, 15), bold_col0=bold_col0)
        if used <= h:
            return used
        remove(sl.shapes[-1])
        fs -= 0.5
    raise DoesNotFit(f"table of {len(data)} rows does not fit {h:.2f}in")


def card(sl: Any, x: float, y: float, w: float, h: float, num: str, title: str, body: str) -> None:
    """A bordered card whose contents are fitted inside the border."""
    rect(sl, x, y, w, h, fill=WHITE, line=RULE, lw=0.9)
    inner = w - 0.40
    top = y + 0.18
    if num:
        tf = txt(sl, x + 0.20, top, inner, 0.24)
        para(tf, num, size=10, color=REFRACT, bold=True, first=True, space_after=0)
        top += 0.28
    used = fitted(
        sl,
        x + 0.20,
        top,
        inner,
        0.72,
        lambda tf, s: para(
            tf,
            title,
            size=s,
            color=HEADING,
            bold=True,
            first=True,
            space_after=0,
            line=1.12,
        ),
        18,
        10.5,
    )
    by = top + used + 0.10
    fitted(
        sl,
        x + 0.20,
        by,
        inner,
        (y + h - 0.14) - by,
        lambda tf, s: para(tf, body, size=s, color=SLATE, first=True, space_after=0, line=1.2),
        16,
        8.5,
    )


def statbar(sl: Any, y: float, stats: list[tuple[str, str]], h: float = 1.35) -> None:
    """Large-number tiles: the figure in crimson, what it counts beneath, both centred."""
    gap = 0.22
    bw = (CW - gap * (len(stats) - 1)) / len(stats)
    for i, (big, label) in enumerate(stats):
        x = ML + i * (bw + gap)
        rect(sl, x, y, bw, h, fill=WHITE, line=RULE, lw=1.0)
        tf = txt(sl, x + 0.15, y + 0.12, bw - 0.30, 0.6, align=PP_ALIGN.CENTER)
        size = 32.0
        while size > 16 and est_lines(big, (bw - 0.30) * SAFETY, size, True, SANS) > 1:
            size -= 1
        para(tf, big, size=size, color=REFRACT, bold=True, first=True, space_after=0, line=1.0)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER

        def write(tf: Any, s: float, lab: str = label) -> None:
            para(tf, lab, size=s, color=SLATE, first=True, space_after=0, line=1.12)
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER

        fitted(sl, x + 0.15, y + 0.74, bw - 0.30, h - 0.82, write, 11.5, 8)


def connect(
    sl: Any, x1: float, y1: float, x2: float, y2: float, color: Any = SLATE, width: float = 1.5
) -> Any:
    """A straight connector. Elbow connectors auto-route into detours that
    collide with the nodes they join; do not use them."""
    c = sl.shapes.add_connector(1, In(x1), In(y1), In(x2), In(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c
