"""
Slide layouts, each drawn from a plain dictionary.

A deck is data: a list of slide specs, each naming its ``kind``. Keeping the
content out of the drawing code is what lets one deck be assembled from three
modules that share one set of layouts, and what keeps every layout short enough to
read, and to fit.

Kinds: ``title``, ``divider``, ``bullets``, ``table``, ``cards``, ``stats``,
``split``, ``flow``, ``context``; and, for the storytelling the deck borrows (a TL;DR
in questions, numbered principles, worked examples, a comparison and a close), ``qa``,
``numbered``, ``workflow``, ``compare`` and ``thanks``.

Copyright (c) 2026 Ashutosh Sinha. All rights reserved.
Proprietary and confidential. See LICENSE and NOTICE at the repository root.
"""

from __future__ import annotations

from typing import Any

import theme as T  # noqa: N812
from pptx.enum.text import PP_ALIGN

from metrics import text_h
from prama.version import VERSION

GAP = 0.18


def _items(tf: Any, items: list[Any], size: float) -> None:
    """Bulleted items; a tuple is ``(head, body)``."""
    for i, it in enumerate(items):
        first = i == 0
        if isinstance(it, tuple):
            T.runs(
                tf,
                [("▪  ", T.REFRACT, True), (it[0], T.INK, True)],
                size=size,
                space_after=1,
                first=first,
                space_before=0 if first else 5,
            )
            T.runs(tf, [(it[1], T.SLATE, False)], size=size - 1.5, space_after=0, level=1)
        else:
            T.runs(
                tf,
                [("▪  ", T.REFRACT, True), (it, T.INK, False)],
                size=size,
                space_after=0,
                first=first,
                space_before=0 if first else 6,
            )


def _intro(sl: Any, y: float, text: str | None) -> float:
    if not text:
        return y
    used = T.fitted(
        sl,
        T.ML,
        y,
        T.CW,
        1.0,
        lambda tf, s: T.para(tf, text, size=s, color=T.INK, first=True, space_after=0, line=1.25),
        14,
        11,
    )
    return y + used + GAP


def _note(sl: Any, text: str | None) -> float:
    """The slide's one-line insight, in a tinted box pinned above the footer; returns its top."""
    if not text:
        return T.BODY_BOTTOM
    size = 13.0
    width = T.CW - 0.5
    while size > 9.5 and text_h(text, width * 0.94, size, False, "Calibri", 1.2) > 0.9:
        size -= 0.5
    h = text_h(text, width * 0.94, size, False, "Calibri", 1.2)
    top = T.BODY_BOTTOM - h - 0.24
    T.rect(sl, T.ML, top, T.CW, h + 0.24, fill=T.TINT)
    T.fitted(
        sl,
        T.ML + 0.25,
        top + 0.11,
        width,
        h + 0.04,
        lambda tf, s: T.para(
            tf, text, size=s, color=T.HEADING, italic=True, first=True, space_after=0, line=1.2
        ),
        size,
        9,
    )
    return top - GAP


def title(s: dict[str, Any]) -> None:
    """The opening slide on the dark gradient: name, promise, author; the deck's map right."""
    T._state["n"] = 1
    sl = T.blank()
    T.ground(sl)
    tf = T.txt(sl, T.ML + 0.1, 1.0, T.CW * 0.6, 0.34)
    T.para(tf, s["kicker"], size=11, color=T.HAZE, bold=True, first=True, space_after=0)

    def head(tf: Any, size: float) -> None:
        for i, line in enumerate(s["title"]):
            T.para(
                tf,
                line,
                size=size,
                color=T.WHITE,
                bold=True,
                first=i == 0,
                space_after=0,
                line=1.05,
            )

    T.fitted(sl, T.ML + 0.1, 1.45, T.CW * 0.58, 1.9, head, 44, 28)
    tf = T.txt(sl, T.ML + 0.1, 3.45, T.CW * 0.58, 0.6)
    T.para(tf, s["sub"], size=24, color=T.HAZE, italic=True, first=True, space_after=0)
    # The prism's six rays: the brand's one flourish, in place of a rule.
    T.spectrum(sl, T.ML + 0.1, 4.25, 1.9, 0.07)
    tf = T.txt(sl, T.ML + 0.1, 4.6, T.CW * 0.5, 1.3)
    T.para(tf, "Ashutosh Sinha", size=20, color=T.WHITE, bold=True, first=True, space_after=3)
    T.para(tf, s.get("date", "October 2026"), size=13, color=T.HAZE, space_after=1)
    T.para(tf, s.get("version", f"Prama {VERSION}"), size=12, color=T.HAZE, space_after=0)
    x0 = T.ML + T.CW * 0.66

    def agenda(tf: Any, size: float) -> None:
        T.para(tf, "IN THIS DECK", size=10, color=T.HAZE, bold=True, first=True, space_after=8)
        for i, c in enumerate(s["agenda"], 1):
            T.runs(
                tf,
                [(f"{i:02d}   ", T.HAZE, True), (c, T.WHITE, False)],
                size=size,
                space_after=4,
            )

    T.fitted(sl, x0, 1.0, T.CW * 0.34, 5.4, agenda, 14, 9)


def divider(s: dict[str, Any]) -> None:
    T.divider(s["num"], s["title"], s["sub"], s["points"])


def bullets(s: dict[str, Any]) -> None:
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    bottom = _note(sl, s.get("note"))
    T.fitted(
        sl,
        T.ML,
        y,
        T.CW,
        bottom - y,
        lambda tf, size: _items(tf, s["items"], size),
        s.get("size", 20),
        9.5,
    )


def table(s: dict[str, Any]) -> None:
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    bottom = _note(sl, s.get("note"))
    T.fitted_table(
        sl,
        s["rows"],
        T.ML,
        y,
        T.CW,
        bottom - y,
        s.get("col_w"),
        start=s.get("size", 17),
        bold_col0=s.get("bold_col0", True),
    )


def cards(s: dict[str, Any]) -> None:
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    bottom = _note(sl, s.get("note"))
    items = s["cards"]
    cols = s.get("cols", 3)
    rows = (len(items) + cols - 1) // cols
    cw = (T.CW - GAP * (cols - 1)) / cols
    ch = (bottom - y - GAP * (rows - 1)) / rows
    for i, (num, head, body) in enumerate(items):
        r, c = divmod(i, cols)
        T.card(sl, T.ML + c * (cw + GAP), y + r * (ch + GAP), cw, ch, num, head, body)


def stats(s: dict[str, Any]) -> None:
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    T.statbar(sl, y, s["stats"])
    y += 1.35 + GAP + 0.1
    bottom = _note(sl, s.get("note"))
    if s.get("rows"):
        T.fitted_table(
            sl, s["rows"], T.ML, y, T.CW, bottom - y, s.get("col_w"), start=s.get("size", 16)
        )
    elif s.get("items"):
        T.fitted(
            sl,
            T.ML,
            y,
            T.CW,
            bottom - y,
            lambda tf, size: _items(tf, s["items"], size),
            s.get("size", 18),
            9.5,
        )


def _column(sl: Any, x: float, y: float, w: float, h: float, col: dict[str, Any]) -> None:
    T.rect(sl, x, y, w, 0.42, fill=T.TINT)
    tf = T.txt(sl, x + 0.18, y + 0.08, w - 0.3, 0.3)
    T.para(tf, col["head"], size=13.5, color=T.HEADING, bold=True, first=True, space_after=0)
    top = y + 0.42 + 0.14
    if col.get("rows"):
        T.fitted_table(
            sl, col["rows"], x, top, w, y + h - top, col.get("col_w"), start=col.get("size", 16)
        )
    else:
        T.fitted(
            sl,
            x,
            top,
            w,
            y + h - top,
            lambda tf, size: _items(tf, col["items"], size),
            col.get("size", 18),
            9,
        )


def split(s: dict[str, Any]) -> None:
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    bottom = _note(sl, s.get("note"))
    w = (T.CW - 0.4) / 2
    _column(sl, T.ML, y, w, bottom - y, s["left"])
    _column(sl, T.ML + w + 0.4, y, w, bottom - y, s["right"])


def flow(s: dict[str, Any]) -> None:
    """Boxes in a row, joined by straight arrows; then optional items beneath."""
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    bottom = _note(sl, s.get("note"))
    steps = s["steps"]
    arrow = 0.28
    bw = (T.CW - arrow * (len(steps) - 1)) / len(steps)
    bh = s.get("box_h", 1.9)
    for i, (head, body) in enumerate(steps):
        x = T.ML + i * (bw + arrow)
        T.card(sl, x, y, bw, bh, "", head, body)
        if i:
            T.connect(sl, x - arrow + 0.03, y + bh / 2, x - 0.03, y + bh / 2, T.REFRACT, 2.0)
    top = y + bh + GAP + 0.05
    if s.get("items"):
        T.fitted(
            sl,
            T.ML,
            top,
            T.CW,
            bottom - top,
            lambda tf, size: _items(tf, s["items"], size),
            s.get("size", 18),
            9,
        )
    elif s.get("rows"):
        T.fitted_table(
            sl, s["rows"], T.ML, top, T.CW, bottom - top, s.get("col_w"), start=s.get("size", 16)
        )


def context(s: dict[str, Any]) -> None:
    """A system context diagram: boxes placed on the content area, joined by arrows.

    Positions are fractions of the content box rather than inches, so a diagram keeps its
    proportions if the theme's margins move. Anchors are chosen from the relative position
    of the two boxes -- an arrow leaves the side that faces its target -- because an elbow
    connector routes itself into the nodes it joins.
    """
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    bottom = _note(sl, s.get("note"))
    x0, w0, h0 = T.ML, T.CW, bottom - y

    box: dict[str, tuple[float, float, float, float]] = {}
    for n in s["nodes"]:
        bx, by = x0 + n["x"] * w0, y + n["y"] * h0
        bw, bh = n["w"] * w0, n["h"] * h0
        box[n["id"]] = (bx, by, bw, bh)

    # Arrows first, so that a box always sits on top of the line that reaches it.
    for e in s.get("edges", []):
        ax, ay, aw, ah = box[e[0]]
        bx, by, bw, bh = box[e[1]]
        acx, acy, bcx, bcy = ax + aw / 2, ay + ah / 2, bx + bw / 2, by + bh / 2
        # Which sides the arrow leaves and enters. Comparing centres is the obvious rule and
        # the wrong one: a wide box can have its centre far to the right of a small box while
        # its left edge is still to the left of it, and the arrow then doubles back on itself.
        # Overlap is the rule that holds -- if the two boxes share a column, the arrow is
        # vertical; if they share a row, it is horizontal -- and a shared column is drawn as
        # a true vertical, down the middle of the overlap rather than slanting between centres.
        #
        # The small gap at each end keeps the connector's bounding box off the card it points
        # at: without it the geometry audit reports a line hidden behind an opaque shape, which
        # from the audit's side is indistinguishable from a line drawn underneath one.
        gap = 0.04
        xlo, xhi = max(ax, bx), min(ax + aw, bx + bw)
        ylo, yhi = max(ay, by), min(ay + ah, by + bh)
        if xhi - xlo > 0.2:
            mid = (xlo + xhi) / 2
            down = bcy > acy
            p1 = (mid, (ay + ah + gap) if down else (ay - gap))
            p2 = (mid, (by - gap) if down else (by + bh + gap))
        elif yhi - ylo > 0.2:
            right = bcx > acx
            mid = (ylo + yhi) / 2
            p1 = ((ax + aw + gap) if right else (ax - gap), mid)
            p2 = ((bx - gap) if right else (bx + bw + gap), mid)
        else:
            down = bcy > acy
            p1 = (acx, (ay + ah + gap) if down else (ay - gap))
            p2 = (bcx, (by - gap) if down else (by + bh + gap))
        T.connect(sl, p1[0], p1[1], p2[0], p2[1], T.SLATE, 1.5)
        # A label only where the arrow has room for one. The gap between two ranks of boxes
        # can be narrower than a line of text, and a label that spills into the card below is
        # worse than no label: the arrow's direction already carries most of the meaning.
        span = abs(p2[1] - p1[1]) if p1[0] == p2[0] else abs(p2[0] - p1[0])
        if len(e) > 2 and e[2] and span > 0.30:
            tf = T.txt(
                sl,
                (p1[0] + p2[0]) / 2 - 0.85,
                (p1[1] + p2[1]) / 2 - 0.085,
                1.7,
                0.17,
                align=PP_ALIGN.CENTER,
            )
            T.para(tf, e[2], size=7.5, color=T.SLATE, space_after=0, first=True)

    for n in s["nodes"]:
        bx, by, bw, bh = box[n["id"]]
        T.card(sl, bx, by, bw, bh, n.get("num", ""), n["head"], n.get("body", ""))


def qa(s: dict[str, Any]) -> None:
    """Question-and-answer rows on tinted bands: the TL;DR. Every answer is set at one size,
    the largest at which all of them fit, so no row reads as more or less important."""
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    bottom = _note(sl, s.get("note"))
    rows = s["rows"]
    gap = 0.14
    rh = (bottom - y - gap * (len(rows) - 1)) / len(rows)
    qw = T.CW * 0.30
    aw = T.CW - qw - 0.35
    size = 15.0
    while size > 9 and any(
        text_h(a, aw * 0.94, size, False, "Calibri", 1.18) > rh - 0.24 for _, a in rows
    ):
        size -= 0.5
    for i, (q, a) in enumerate(rows):
        top = y + i * (rh + gap)
        T.rect(sl, T.ML, top, T.CW, rh, fill=T.TINT)
        T.fitted(
            sl,
            T.ML + 0.25,
            top + 0.12,
            qw - 0.35,
            rh - 0.24,
            lambda tf, sz, q=q: T.para(
                tf, q, size=sz, color=T.HEADING, bold=True, first=True, space_after=0, line=1.1
            ),
            16,
            10,
        )
        T.fitted(
            sl,
            T.ML + qw + 0.1,
            top + 0.12,
            aw,
            rh - 0.24,
            lambda tf, sz, a=a: T.para(
                tf, a, size=sz, color=T.INK, first=True, space_after=0, line=1.18
            ),
            size,
            9,
        )


def _numbered_cell(
    sl: Any, x: float, y: float, w: float, h: float, n: int, head: str, body: str
) -> None:
    d = 0.46
    T.circle(sl, x, y + 0.02, d, str(n), size=14)

    def write(tf: Any, size: float) -> None:
        T.para(tf, head, size=size, color=T.HEADING, bold=True, first=True, space_after=2, line=1.1)
        if body:
            T.para(tf, body, size=size - 2.5, color=T.SLATE, space_after=0, line=1.15)

    T.fitted(sl, x + d + 0.2, y, w - d - 0.2, h, write, 16, 10.5)


def numbered(s: dict[str, Any]) -> None:
    """Numbered circles, each with a bold head and a line beneath: principles, steps."""
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    bottom = _note(sl, s.get("note"))
    items = s["items"]
    cols = s.get("cols", 1 if len(items) <= 5 else 2)
    per = (len(items) + cols - 1) // cols
    cw = (T.CW - 0.45 * (cols - 1)) / cols
    rh = (bottom - y) / per
    for i, (head, body) in enumerate(items):
        c, r = divmod(i, per)
        _numbered_cell(sl, T.ML + c * (cw + 0.45), y + r * rh, cw, rh - 0.08, i + 1, head, body)


def workflow(s: dict[str, Any]) -> None:
    """A worked example: the business question in quotes, then each step, numbered,
    with what runs and what comes back; the insight below."""
    sl, y = T.content(s["title"], s.get("kicker"))
    q = f"“{s['question']}”"
    used = T.fitted(
        sl,
        T.ML,
        y,
        T.CW,
        0.8,
        lambda tf, size: T.para(
            tf, q, size=size, color=T.REFRACT, italic=True, first=True, space_after=0, line=1.2
        ),
        18,
        12,
    )
    y += used + GAP + 0.06
    bottom = _note(sl, s.get("note"))
    steps = s["steps"]
    rh = (bottom - y) / len(steps)
    nw = T.CW * 0.30
    d = 0.42
    for i, (name, detail) in enumerate(steps):
        top = y + i * rh
        T.circle(sl, T.ML, top + 0.02, d, str(i + 1), size=13)
        T.fitted(
            sl,
            T.ML + d + 0.2,
            top + 0.04,
            nw - d - 0.2,
            rh - 0.1,
            lambda tf, size, n=name: T.para(
                tf,
                n,
                size=size,
                color=T.HEADING,
                bold=True,
                first=True,
                space_after=0,
                line=1.1,
            ),
            18,
            10,
        )
        T.fitted(
            sl,
            T.ML + nw + 0.1,
            top + 0.04,
            T.CW - nw - 0.1,
            rh - 0.1,
            lambda tf, size, dt=detail: T.para(
                tf, dt, size=size, color=T.INK, first=True, space_after=0, line=1.15
            ),
            17,
            9.5,
        )


def compare(s: dict[str, Any]) -> None:
    """A comparison table whose verdict words are coloured: yes, partly, no."""
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    bottom = _note(sl, s.get("note"))
    T.fitted_table(
        sl, s["rows"], T.ML, y, T.CW, bottom - y, s.get("col_w"), start=s.get("size", 14)
    )
    tbl = sl.shapes[-1].table
    marks = {"Yes": T.OK, "Partly": T.WARN, "No": T.REFRACT}
    for r in range(1, len(s["rows"])):
        for c in range(1, len(s["rows"][0])):
            cell = tbl.cell(r, c)
            first = cell.text.split(" ")[0].rstrip(":,;")
            if first in marks:
                for run in cell.text_frame.paragraphs[0].runs:
                    run.font.color.rgb = marks[first]
                    run.font.bold = True


def thanks(s: dict[str, Any]) -> None:
    """The close, on the dark gradient: the promise once more, and who to ask."""
    T._state["n"] += 1
    sl = T.blank()
    T.ground(sl)
    tf = T.txt(sl, T.ML + 0.1, 2.0, T.CW * 0.8, 1.0)
    T.para(
        tf,
        s.get("title", "Thank you"),
        size=48,
        color=T.WHITE,
        bold=True,
        first=True,
        space_after=0,
    )
    tf = T.txt(sl, T.ML + 0.1, 3.15, T.CW * 0.8, 0.6)
    T.para(tf, s["sub"], size=24, color=T.HAZE, italic=True, first=True, space_after=0)
    T.spectrum(sl, T.ML + 0.1, 3.95, 1.9, 0.07)
    tf = T.txt(sl, T.ML + 0.1, 4.3, T.CW * 0.8, 1.6)
    for i, line in enumerate(s["lines"]):
        T.para(
            tf,
            line,
            size=16 if i else 20,
            color=T.WHITE if i == 0 else T.HAZE,
            bold=i == 0,
            first=i == 0,
            space_after=4,
        )


KINDS = {
    "qa": qa,
    "numbered": numbered,
    "workflow": workflow,
    "compare": compare,
    "thanks": thanks,
    "context": context,
    "title": title,
    "divider": divider,
    "bullets": bullets,
    "table": table,
    "cards": cards,
    "stats": stats,
    "split": split,
    "flow": flow,
}


def render(slides: list[dict[str, Any]]) -> None:
    for i, spec in enumerate(slides, 1):
        try:
            KINDS[spec["kind"]](spec)
        except T.DoesNotFit as exc:
            raise T.DoesNotFit(f"slide {i} ({spec.get('title')!r}): {exc}") from exc
